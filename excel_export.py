"""
excel_export.py
================
Builds a professional, formatted Excel (.xlsx) historical data report
using openpyxl. This is the one additional Python module referenced
in app.py's /api/export/excel route.
"""

from io import BytesIO
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


HEADER_FILL = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
TITLE_FONT = Font(bold=True, size=14, color="0F172A")
THIN_BORDER = Border(
    left=Side(style="thin", color="D1D5DB"),
    right=Side(style="thin", color="D1D5DB"),
    top=Side(style="thin", color="D1D5DB"),
    bottom=Side(style="thin", color="D1D5DB"),
)

COLUMNS = [
    ("Timestamp", "timestamp", 20),
    ("T1 - Comp. Suction (°C)", "T1", 20),
    ("T2 - Comp. Discharge (°C)", "T2", 22),
    ("T3 - Condenser Outlet (°C)", "T3", 22),
    ("T4 - Evaporator Inlet (°C)", "T4", 22),
    ("T5 - Water Inlet (°C)", "T5", 18),
    ("T6 - Water Outlet (°C)", "T6", 18),
    ("Flow Rate (L/min)", "flow_rate", 16),
    ("Power (kW)", "power_kw", 12),
    ("Cooling Capacity (kW)", "cooling_capacity", 18),
    ("Water ΔT (°C)", "delta_t", 14),
    ("COP", "cop", 10),
    ("Compressor Status", "compressor_status", 16),
]


def _fmt(value):
    return value if value is not None else "N/A"


def build_excel_report(rows, cfg):
    """
    rows: list of sqlite row dicts (already filtered to date range)
    cfg:  the config module (for machine info in the summary sheet)
    Returns: raw xlsx bytes
    """
    wb = Workbook()

    # ---------------- DATA SHEET ----------------
    ws = wb.active
    ws.title = "Historical Data"

    ws.merge_cells("A1:M1")
    ws["A1"] = "Chiller Historical Data Report"
    ws["A1"].font = TITLE_FONT
    ws.merge_cells("A2:M2")
    ws["A2"] = f"Machine: {cfg.MACHINE_INFO['manufacturer']} {cfg.MACHINE_INFO['model']}  |  Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
    ws["A2"].font = Font(italic=True, size=9, color="6B7280")

    header_row = 4
    for i, (label, key, width) in enumerate(COLUMNS, start=1):
        c = ws.cell(row=header_row, column=i, value=label)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = THIN_BORDER
        ws.column_dimensions[get_column_letter(i)].width = width

    r = header_row + 1
    for row in rows:
        for i, (label, key, width) in enumerate(COLUMNS, start=1):
            if key == "compressor_status":
                val = "RUNNING" if row.get(key) else "STOPPED"
            elif key == "timestamp":
                val = row.get(key)
            else:
                val = _fmt(row.get(key))
            c = ws.cell(row=r, column=i, value=val)
            c.border = THIN_BORDER
            c.alignment = Alignment(horizontal="center")
        r += 1

    # Freeze header row, add autofilter / excel table styling
    ws.freeze_panes = f"A{header_row + 1}"
    last_col_letter = get_column_letter(len(COLUMNS))
    if r > header_row + 1:
        ws.auto_filter.ref = f"A{header_row}:{last_col_letter}{r - 1}"

    # ---------------- SUMMARY SHEET ----------------
    ws2 = wb.create_sheet("Summary")
    ws2.column_dimensions["A"].width = 34
    ws2.column_dimensions["B"].width = 18

    ws2.merge_cells("A1:B1")
    ws2["A1"] = "Report Summary"
    ws2["A1"].font = TITLE_FONT

    cops = [row["cop"] for row in rows if row.get("cop") is not None]
    caps = [row["cooling_capacity"] for row in rows if row.get("cooling_capacity") is not None]
    powers = [row["power_kw"] for row in rows if row.get("power_kw") is not None]
    t2s = [row["T2"] for row in rows if row.get("T2") is not None]
    dts = [row["delta_t"] for row in rows if row.get("delta_t") is not None]

    def avg(lst):
        return round(sum(lst) / len(lst), 3) if lst else "N/A"

    def mx(lst):
        return round(max(lst), 3) if lst else "N/A"

    def mn(lst):
        return round(min(lst), 3) if lst else "N/A"

    summary_items = [
        ("Average COP", avg(cops)),
        ("Maximum COP", mx(cops)),
        ("Minimum COP", mn(cops)),
        ("Average Cooling Capacity (kW)", avg(caps)),
        ("Maximum Cooling Capacity (kW)", mx(caps)),
        ("Average Power (kW)", avg(powers)),
        ("Maximum Compressor Discharge T2 (°C)", mx(t2s)),
        ("Average Water ΔT (°C)", avg(dts)),
        ("Number of Alarm-Triggering Records", sum(
            1 for row in rows if (row.get("T2") and row["T2"] > cfg.ALARM_THRESHOLDS["T2_WARNING"])
        )),
        ("Total Operating Records", len(rows)),
    ]

    row_i = 3
    for label, value in summary_items:
        lc = ws2.cell(row=row_i, column=1, value=label)
        lc.font = Font(bold=True)
        lc.border = THIN_BORDER
        vc = ws2.cell(row=row_i, column=2, value=value)
        vc.border = THIN_BORDER
        vc.alignment = Alignment(horizontal="center")
        row_i += 1

    # Machine info block
    row_i += 1
    ws2.cell(row=row_i, column=1, value="Machine Information").font = Font(bold=True, size=12)
    row_i += 1
    for k, v in cfg.MACHINE_INFO.items():
        ws2.cell(row=row_i, column=1, value=k.replace("_", " ").title()).font = Font(bold=True)
        ws2.cell(row=row_i, column=2, value=v)
        row_i += 1

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
